"""
L0 输出物导出（Word）

这一层的风险集中在两处，都不是功能性的：

1. **中文显示**——字体没设 eastAsia 的话，Word 里中文掉回默认字体而西文正常，
   中英混排的备忘录会显示成两种字号拼接，一眼就看出来。
2. **中文文件名**——Content-Disposition 没走 RFC 5987 的话，下载下来叫
   %E5%86%B3%E7%AD%96…docx，演示现场当场社死。
"""
import io
import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault("USE_MOCK", "True")

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from core.docx_export import markdown_to_docx, markdown_to_docx_bytes
from routers.report import _content_disposition

EAST_ASIA = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"

SAMPLE = """\
# 主诉评估结果：测试案

> 案由：商标侵权 ｜ 业务目标：要钱 ｜ 评估模型：v4 二维主诉决策模型

## 一、核心结论

- **主诉决策分**：69.3（法律可行性 68.7 与 业务预期 70.0 的均衡水平）
- **决策象限**：双低区（暂缓，评估替代方案）
  - 缩进项

## 二、维度明细

### 红线检查

- ⛔ 诉讼时效检查：案情描述中未明确侵权行为发现时间
- ✅ 主体资格检查：已录入原告，且持有权利基础证明

### 权利基础（82 分）

商标合法注册且在有效期内。

1. 有序项一
2. 有序项二

---

> 本评估结果由 Soft IP 主诉评估系统生成。
"""


@pytest.fixture
def docx_bytes():
    return markdown_to_docx_bytes(SAMPLE, meta=["生成日期：2026-09-01 ｜ 当前状态（未定稿）"])


def _paragraphs(data: bytes):
    return [p for p in Document(io.BytesIO(data)).paragraphs if p.text.strip()]


# ============================================================
# A. 文件本身
# ============================================================

class TestFileFormat:

    def test_output_is_a_real_docx(self, docx_bytes):
        """能被引擎打开才算数，否则用户下载到的是个打不开的文件"""
        assert docx_bytes[:2] == b"PK"                 # zip 容器
        assert zipfile.is_zipfile(io.BytesIO(docx_bytes))
        Document(io.BytesIO(docx_bytes))               # 解析失败会抛

    def test_meta_line_follows_title(self, docx_bytes):
        """标题在上、生成日期紧随其后：层级正挂，日期不能压在主标题上面"""
        paragraphs = _paragraphs(docx_bytes)
        assert paragraphs[0].text.startswith("主诉评估结果")
        assert paragraphs[1].text.startswith("生成日期：")


# ============================================================
# B. 排版映射
# ============================================================

class TestMarkdownMapping:

    def test_heading_levels_are_distinct(self, docx_bytes):
        """三级标题字号必须递减，否则 Word 里分不出层级"""

        def size_of(prefix):
            hit = next(p for p in _paragraphs(docx_bytes) if p.text.strip().startswith(prefix))
            return hit.runs[0].font.size.pt

        assert size_of("主诉评估结果") > size_of("一、") > size_of("权利基础（")

    def test_inline_bold_is_parsed(self, docx_bytes):
        """
        **加粗** 必须真的变成加粗，而不是把星号原样印在 Word 里。

        星号漏出去是最显眼的一类排版事故——收件人一眼就看出是程序生成的。
        """
        target = next(p for p in _paragraphs(docx_bytes)
                      if "主诉决策分" in p.text)
        assert "**" not in target.text
        bold_runs = [r.text for r in target.runs if r.bold]
        assert any("主诉决策分" in t for t in bold_runs)

    def test_bullets_become_list_paragraphs(self, docx_bytes):
        bullets = [p for p in Document(io.BytesIO(docx_bytes)).paragraphs
                   if p.style.name == "List Bullet"]
        assert len(bullets) >= 3

    def test_blockquote_is_not_a_plain_paragraph(self, docx_bytes):
        quote = next(p for p in _paragraphs(docx_bytes) if "评估模型" in p.text)
        assert quote.paragraph_format.left_indent is not None

    def test_horizontal_rule_adds_border(self, docx_bytes):
        """分隔线靠段落下边框实现，而不是塞一串破折号"""
        doc = Document(io.BytesIO(docx_bytes))
        assert any("pBdr" in p._p.xml for p in doc.paragraphs)

    def test_emoji_and_symbols_survive(self, docx_bytes):
        """⛔⚠️✅ 这些标记在报告里大量出现，不能丢"""
        text = "\n".join(p.text for p in _paragraphs(docx_bytes))
        assert "✅" in text

    def test_empty_markdown_does_not_crash(self):
        assert markdown_to_docx_bytes("")[:2] == b"PK"


TABLE_MD = """\
# 主诉评估结果：表格测试案

## 一、维度明细

| 评估维度 | 得分 |
| --- | ---: |
| 权利基础 | 78.0 |

## 二、证据盘点与缺口清单

| 缺口项 | 类别 | 支撑要件 | 补证建议 |
| --- | --- | --- | --- |
| 商标注册证原件 | 权利基础 | 权利有效性 | 需补充续展证明 |
"""


class TestTableRendering:
    """维度汇总表 / 证据缺口表是 GFM 管道表格，必须渲染成真表格而非裸竖线"""

    def _doc(self) -> Document:
        return Document(io.BytesIO(markdown_to_docx_bytes(TABLE_MD)))

    def test_tables_are_real_tables(self):
        assert len(self._doc().tables) == 2, "两张 GFM 表应渲染成两个 Word 表格"

    def test_header_row_is_dark_blue_with_white_text(self):
        cell = self._doc().tables[0].rows[0].cells[0]
        assert cell.text == "评估维度"
        shd = cell._tc.tcPr.find(qn("w:shd"))
        assert shd is not None and shd.get(qn("w:fill")) == "1F4E79"
        run = cell.paragraphs[0].runs[0]
        assert run.font.color.rgb == RGBColor(0xFF, 0xFF, 0xFF)

    def test_header_row_repeats_across_pages(self):
        """跨页表格必须重复表头，否则第二页的裸数据行没有列含义"""
        for table in self._doc().tables:
            tr_pr = table.rows[0]._tr.trPr
            assert tr_pr is not None and tr_pr.find(qn("w:tblHeader")) is not None

    def test_no_raw_pipe_leaks_into_paragraphs(self):
        """竖线是表格语法字符，漏进正文段落就说明表格没被识别"""
        for p in self._doc().paragraphs:
            assert "|" not in p.text

    def test_data_rows_carry_content(self):
        table = self._doc().tables[0]
        assert table.rows[1].cells[0].text == "权利基础"
        assert table.rows[1].cells[1].text == "78.0"


# ============================================================
# C. 中文字体（最容易漏的一项）
# ============================================================

class TestChineseTypography:

    def test_east_asia_font_is_set(self, docx_bytes):
        """
        只设 font.name 的话 Word 里中文会掉回默认字体，而西文是设好的——
        中英混排会显示成两种字号的拼接。必须显式设 eastAsia。
        """
        for p in _paragraphs(docx_bytes):
            for run in p.runs:
                fonts = run._element.rPr.rFonts
                assert fonts.get(EAST_ASIA) == "宋体", f"未设中文字体：{p.text[:20]}"

    def test_body_size_is_five_point(self, docx_bytes):
        """10.5pt（五号）是中文正式文档通例，比 Word 默认 11pt 更紧凑"""
        body = next(p for p in _paragraphs(docx_bytes) if "商标合法注册" in p.text)
        assert body.runs[0].font.size.pt == 10.5


# ============================================================
# D. 中文文件名
# ============================================================

class TestContentDisposition:

    def test_rfc5987_and_ascii_fallback_both_present(self):
        from urllib.parse import unquote

        header = _content_disposition("某某商标案-评估结果.docx")
        assert "filename*=UTF-8''" in header
        assert 'filename="' in header
        # 中文名确实被编码进去了，且能原样解回来
        encoded = header.split("filename*=UTF-8''")[1]
        assert unquote(encoded) == "某某商标案-评估结果.docx"

    def test_ascii_fallback_drops_rather_than_mangles(self):
        """
        非 ASCII 字符应当**丢掉**而不是替换成下划线。

        「E2E-版本-决策备忘录」替换成下划线会变成 E2E-__-_____ 这种
        谁都不想看到的东西，丢掉则是干净的 E2E.docx。
        """
        assert _content_disposition("E2E-版本-评估结果.docx").startswith(
            'attachment; filename="E2E.docx"')

    def test_pure_chinese_name_still_gets_a_usable_fallback(self):
        header = _content_disposition("评估结果.docx")
        assert 'filename="report.docx"' in header

    def test_illegal_filename_chars_are_stripped(self):
        """引号会截断 header，斜杠会破坏路径——都必须清掉"""
        header = _content_disposition('a/b:c*d?e"f.docx')
        assert 'filename="abcdef.docx"' in header


# ============================================================
# E. 接口
# ============================================================

class TestExportEndpoints:
    """导出接口走真实 HTTP，覆盖版本选择与缺失数据的分支"""

    DOCX_MAGIC = b"PK"

    @pytest.fixture
    def case_id(self):
        from core.database import Case, Report, SessionLocal, init_db
        init_db()
        db = SessionLocal()
        try:
            # 每个用例后清干净：固定 id 便于排查，但重复插入会撞主键
            db.query(Report).filter(Report.case_id == "docx-case").delete()
            db.query(Case).filter(Case.id == "docx-case").delete()
            db.add(Case(id="docx-case", name="某某商标案", cause_type="商标侵权",
                        goal_type="要钱", status="completed", context_json={}))
            db.commit()
        finally:
            db.close()
        yield "docx-case"

        db = SessionLocal()
        try:
            db.query(Report).filter(Report.case_id == "docx-case").delete()
            db.query(Case).filter(Case.id == "docx-case").delete()
            db.commit()
        finally:
            db.close()

    @pytest.fixture
    def client(self, case_id):
        from fastapi.testclient import TestClient
        from main import app
        with TestClient(app) as c:
            yield c

    def test_result_export_returns_docx(self, client, case_id):
        resp = client.get(f"/api/report/{case_id}/result.docx")
        assert resp.status_code == 200
        assert resp.content[:2] == self.DOCX_MAGIC
        assert "wordprocessingml" in resp.headers["content-type"]

    def test_result_header_carries_utf8_filename(self, client, case_id):
        header = client.get(f"/api/report/{case_id}/result.docx").headers["content-disposition"]
        assert "filename*=UTF-8''" in header

    def test_transcript_export_requires_a_moot(self, client, case_id):
        """没跑过庭审就导出，应当是明确 404 而不是一份空文档"""
        assert client.get(f"/api/report/{case_id}/transcript.docx").status_code == 404

    def test_transcript_export_when_moot_exists(self, client, case_id):
        from core.case_context import CaseContext
        from core.database import SessionLocal
        db = SessionLocal()
        try:
            from routers.report import _load_case
            case = _load_case(db, case_id)
            ctx = CaseContext.from_dict({})
            ctx.case_id = case_id
            ctx.cause_type, ctx.goal_type = "商标侵权", "要钱"
            ctx.moot_transcript = [
                {"step": 1, "step_name": "原告陈述", "role": "plaintiff",
                 "role_name": "原告代理人", "content": "我方商标合法注册且在有效期内。"},
            ]
            ctx.correction_coeff = 0.9
            case.context_json = ctx.to_dict()
            db.commit()
        finally:
            db.close()

        resp = client.get(f"/api/report/{case_id}/transcript.docx")
        assert resp.status_code == 200
        text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
        assert "原告陈述" in text and "0.9" in text
